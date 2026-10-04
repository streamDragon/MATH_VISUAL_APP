#!/usr/bin/env python3
"""Extract review drafts from supplied TXT, HTML, DOCX or text-based PDF.

No AI, auto-publication or arbitrary web crawler. All candidates need review.
"""
import argparse
import hashlib
import json
import re
from html.parser import HTMLParser
from pathlib import Path


class PageText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript'):
            self.hidden += 1
        if tag in ('p', 'div', 'li', 'br', 'h1', 'h2', 'h3', 'tr'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript'):
            self.hidden = max(0, self.hidden - 1)
        if tag in ('p', 'div', 'li', 'h1', 'h2', 'h3', 'tr'):
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def extract(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        return [(i + 1, page.extract_text() or '') for i, page in enumerate(PdfReader(path).pages)]
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        return [(None, text)]
    text = path.read_text(encoding='utf-8-sig')
    if suffix in ('.html', '.htm'):
        parser = PageText()
        parser.feed(text)
        text = ''.join(parser.parts)
    elif suffix not in ('.txt', '.md'):
        raise ValueError('Supported inputs: .pdf, .docx, .txt, .md, .html')
    return [(None, text)]


def candidates(text):
    # Numbered questions are a heuristic, not a claim of mathematical parsing.
    text = text.replace('\r', '')
    marker = r'(?m)^\s*(?:שאלה\s+\d+|Question\s+\d+|\d{1,3}[.)])\s*'
    starts = list(re.finditer(marker, text, re.IGNORECASE))
    if starts:
        return [text[m.start():starts[i + 1].start() if i + 1 < len(starts) else len(text)].strip()
                for i, m in enumerate(starts)]
    return [part.strip() for part in re.split(r'\n\s*\n', text) if part.strip()]


def prepare(pages, source, grade, language, permission_note):
    result, seen = [], set()
    for page, text in pages:
        for question in candidates(text):
            digest = hashlib.sha256(question.encode()).hexdigest()
            if len(question) < 12 or digest in seen:
                continue
            seen.add(digest)
            result.append({'id': 'draft-' + digest[:16], 'status': 'needs_human_review',
                           'question_text': question, 'grade': grade, 'language': language,
                           'source': source, 'source_page': page, 'permission_note': permission_note,
                           'extraction': 'heuristic_text', 'activity': None})
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--source', required=True, help='Book citation or original page URL')
    parser.add_argument('--permission-note', required=True, help='Why this material may be used')
    parser.add_argument('--grade', type=int, choices=(9, 10, 11), default=9)
    parser.add_argument('--language', choices=('he', 'en'), default='he')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        rows = prepare(extract(args.input), args.source, args.grade, args.language, args.permission_note)
        if not rows:
            raise ValueError('No usable text found. Scanned pages require OCR and human correction first.')
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f'{len(rows)} review drafts saved. Nothing was published; formulas and boundaries need review.')
    except (ValueError, ImportError) as error:
        parser.exit(1, str(error) + '\n')


if __name__ == '__main__':
    main()
